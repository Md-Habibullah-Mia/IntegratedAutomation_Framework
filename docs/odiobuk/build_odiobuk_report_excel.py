"""Builds the Odiobuk test report workbook (web + mobile), like the LexVerify one.

  python docs/odiobuk/build_odiobuk_report_excel.py <full-run.json> [rerun.json ...]

Web rows come from Playwright JSON reports: the full odiobuk-chromium run
(including @slow), then any re-runs, whose results replace the earlier ones
for the tests they contain; mobile rows are the results of the 2026-09-30 emulator
run (see CHANGELOG Phase 18), with each test's expected result looked up in
the mobile test-case matrix (docs/MemoryWave_Test_Cases.xlsx) by TC ID.
"""
import json
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DOCS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WEB_RUN_DATE = "2026-10-02"
MOBILE_RUN_DATE = "2026-09-30"
WEB_ENV = "https://odiobuk.eastus2.cloudapp.azure.com (chromium, 1 worker)"
MOBILE_ENV = "Android emulator Pixel 7 · API 34 (google_apis) · memorywave_v2.apk (release 'mobile-apk', 2026-09-22)"

PASS, FAIL, SKIP, BLOCKED = "Pass", "Fail", "Skipped", "Blocked"

# ------------------------------------------------------------------ Web
WEB_MODULES = {
    "registration.spec.ts": "Registration",
    "login.spec.ts": "Login",
    "home.spec.ts": "Home",
    "library.spec.ts": "Library",
    "saved.spec.ts": "Saved",
    "voices.spec.ts": "Voices",
    "documents.spec.ts": "Admin · PDF catalogue",
    "profile.spec.ts": "Profile (You)",
    "narration-e2e.spec.ts": "E2E · PDF narration",
    "narration-catalogue-e2e.spec.ts": "E2E · Catalogue narration",
    "voice-privacy-security.spec.ts": "Security",
}
# Notes added after reading the run (keyed by TC ID); failures get their
# error text automatically.
ADMIN_BLOCKED = ("Blocked: the admin test account (MOBILE_TEST_ADMIN_EMAIL) is rejected — UI \"Invalid email or "
                 "password\", POST /api/auth/login → 401 (checked 2026-10-02). Needs the current admin credentials.")
NETWORK_RERUN = ("First run hit a laptop network error (DNS ERR_NAME_NOT_RESOLVED / ERR_ABORTED) at the same moment as "
                 "two other tests; re-run on its own and passed.")
WEB_NOTES = {
    "TC-001": "Sign-up now lands on a new /welcome step (\"What brings you here?\"); tests skip it as a user can.",
    "TC-007": NETWORK_RERUN,
    "TC-008": NETWORK_RERUN,
    "TC-009": "Did not run in the full run (serial file stopped at TC-008's network error); passed on re-run.",
    "TC-012": "Test updated: deleting now also requires the account password; \"DELETE\" alone keeps the button disabled.",
    "TC-015": "Real narration of a free catalogue title with a house voice — finished in ~5.6 min.",
    "SEC-001": ("Re-targeted to the current voice (HH2) and a brand-new unrelated account (admin login is broken). "
                "Unrelated account now gets HTTP 403 — the voice-preview authorization bug (issue #13) is FIXED."),
}
# Tests whose outcome is decided by something outside the test (status, note).
WEB_OVERRIDES = {tc: (BLOCKED, ADMIN_BLOCKED) for tc in ("TC-003", "TC-010", "TC-014", "TC-013")}


def web_rows(report_paths):
    by_id = {}
    for path in report_paths:
        for row in _web_rows(path):
            by_id[row["id"]] = row  # later reports (re-runs) win
    for tc, (status, note) in WEB_OVERRIDES.items():
        if tc in by_id:
            by_id[tc].update(status=status, note=note)
    order = lambda r: (r["id"].startswith("SEC"), int(re.sub(r"\D", "", r["id"]) or 0), r["id"])
    return sorted(by_id.values(), key=order)


def _web_rows(report_path):
    with open(report_path, encoding="utf-8") as f:
        report = json.load(f)
    rows = []

    def walk(suite, describe):
        for spec in suite.get("specs", []):
            for t in spec.get("tests", []):
                results = t.get("results") or [{}]
                last = results[-1]
                outcome = t.get("status")  # expected | unexpected | flaky | skipped
                status = {"expected": PASS, "flaky": PASS, "unexpected": FAIL, "skipped": SKIP}.get(outcome, outcome)
                title = spec["title"]
                tc = title.split(" - ")[0].strip()
                file = os.path.basename(spec.get("file", ""))
                err = (last.get("error") or {}).get("message", "")
                err = re.sub(r"\x1b\[[0-9;]*m", "", err).strip().splitlines()
                note = WEB_NOTES.get(tc, "")
                if status == FAIL and err:
                    note = (note + " " if note else "") + " | ".join(l.strip() for l in err[:4] if l.strip())
                if outcome == "flaky":
                    note = (note + " " if note else "") + f"Flaky: passed on retry {len(results) - 1}."
                rows.append({
                    "id": tc,
                    "module": WEB_MODULES.get(file, file),
                    "title": title.split(" - ", 1)[-1].replace(" @slow", ""),
                    "type": "E2E (@slow)" if "@slow" in title else describe.split(" - ")[0],
                    "spec": file,
                    "status": status,
                    "seconds": round(sum(r.get("duration", 0) for r in results) / 1000, 1),
                    "note": note,
                })
        for child in suite.get("suites", []):
            walk(child, child.get("title", describe))

    for s in report.get("suites", []):
        walk(s, s.get("title", ""))
    return rows


# ------------------------------------------------------------------ Mobile
# (spec, test title, status, note) — from the 2026-09-30 emulator run.
MOBILE = [
    ("login.spec.ts", "Shows an error with invalid credentials", PASS, ""),
    ("login.spec.ts", "Shows an error when email is left empty", PASS, ""),
    ("login.spec.ts", "Shows an error when password is left empty", PASS, ""),
    ("login.spec.ts", "Rejects a malformed email with client-side validation", PASS, ""),
    ("login.spec.ts", "Rejects a SQL-injection-shaped email with the same client-side validation", PASS, ""),
    ("login.spec.ts", "Logs in successfully with valid credentials", PASS, ""),
    ("registration.spec.ts", "Shows required-field errors when submitted empty", PASS, ""),
    ("registration.spec.ts", "Rejects a password under 8 characters", PASS, ""),
    ("registration.spec.ts", "Rejects mismatched password and confirmation", PASS, ""),
    ("registration.spec.ts", "Rejects an email that is already registered", PASS, ""),
    ("setup/verified-login.spec.ts", "LGN-010: a pre-verified account logs in straight to Home, skipping verification", PASS,
     "First-Login Verification flag seeded locally (no server-side equivalent), then one real login."),
    ("home.spec.ts", "HOM-003: shows a time-based greeting with the account name", PASS, ""),
    ("home.spec.ts", "HOM-011a: lists the account's real audiobooks under Your Audiobooks", PASS,
     "Re-enabled 2026-09-30 — \"And Then There Were None\"."),
    ("home.spec.ts", "HOM-011b / AUD-034: tapping a playable audiobook opens the player with the correct title", PASS, "Re-enabled 2026-09-30."),
    ("home.spec.ts", "HOM-013a: lists the account's real PDF under Your Books · PDF", SKIP,
     "Blocked: own-PDF upload is disabled server-side (POST /api/documents → 403), so the account can't have a PDF."),
    ("home.spec.ts", "HOM-013b / LIB-023 / LIB-024: tapping the PDF opens Book Pages with the right filename and page list", SKIP,
     "Blocked: same as HOM-013a."),
    ("home.spec.ts", "HOM-016: bottom navigation switches tabs and each shows its own real header", PASS, ""),
    ("library.spec.ts", "LIB-002a: Public shelf header shows a real title count", PASS, ""),
    ("library.spec.ts", "LIB-002b / LIB-003: Your Library shelf shows a real item count and its own dashed actions", PASS, ""),
    ("library.spec.ts", "LIB-004: category chips switch which rows are shown on the Public shelf", PASS,
     "Test updated: checks row shapes instead of the old sample titles."),
    ("library.spec.ts", "LIB-011: tapping an already-owned voice shows a toast, not the acquire dialog", PASS, "Re-enabled 2026-09-30 — \"Sumitra\"."),
    ("audiobooks.spec.ts", "AUD-004/AUD-015: opening an owned catalogue title pre-fills an existing voice, title, and the store-source note", PASS,
     "Re-enabled 2026-09-30."),
    ("audiobooks.spec.ts", "AUD-021: Generate produces the real \"Narration started\" confirmation, using the existing voice", PASS,
     "Real generation with voice HH2; new row reached DONE in ~1 min."),
    ("audiobooks.spec.ts", "AUD-034/AUD-036: opening a finished audiobook from the history list opens its player with the right title", PASS,
     "Re-enabled 2026-09-30."),
    ("voices.spec.ts", "VOI header, tagline, and record CTA render correctly", PASS, ""),
    ("voices.spec.ts", "VOI-009: Your Voices lists an unregistered capture session", PASS, ""),
    ("voices.spec.ts", "VOI-004: Your Voices lists a real ready voice", PASS, "Re-enabled 2026-09-30 — \"HH2\"."),
    ("voices.spec.ts", "VOI-010/VOI-011: Register dialog opens pre-filled and Cancel discards without mutating", PASS, ""),
    ("voices.spec.ts", "VOI-016/VOI-017: Lend dialog shows the real consent text and \"Not now\" declines without mutating", PASS,
     "Re-enabled 2026-09-30."),
    ("profile.spec.ts", "PROF-001: Profile loads with real identity, stats, usage, and premium cards", PASS, ""),
    ("profile.spec.ts", "PROF-028: Account section shows real sign-in method and membership info", PASS, ""),
    ("profile.spec.ts", "PROF-019: declining the Upgrade dialog makes no plan change", PASS, ""),
    ("profile.spec.ts", "PROF-029: \"Private vault\" navigates to the real Vault screen", PASS, ""),
    ("vault.spec.ts", "VLT-001/VLT-003: Vault loads and shows the real encryption/verification status", PASS, ""),
    ("vault.spec.ts", "VLT-004: empty vault shows the real empty state and import CTA", PASS, ""),
]
MOBILE_MODULES = {
    "login.spec.ts": "Login", "registration.spec.ts": "Registration", "setup/verified-login.spec.ts": "Verified session",
    "home.spec.ts": "Home", "library.spec.ts": "Library", "audiobooks.spec.ts": "Audiobooks & Player",
    "voices.spec.ts": "Voices", "profile.spec.ts": "Profile", "vault.spec.ts": "Vault",
}


def matrix_expected():
    """TC ID → (title, expected) from the mobile test-case matrix."""
    out = {}
    path = os.path.join(DOCS, "MemoryWave_Test_Cases.xlsx")
    if not os.path.exists(path):
        return out
    wb = load_workbook(path, read_only=True)
    for name in wb.sheetnames[1:]:
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                out[str(row[0]).strip()] = (row[1], row[4])
    return out


def mobile_rows():
    matrix = matrix_expected()
    rows = []
    for spec, title, status, note in MOBILE:
        ids = re.findall(r"\b([A-Z]{3,4}-\d{3}[a-z]?)", title.split(":")[0])
        expected = ""
        for i in ids:
            hit = matrix.get(i) or matrix.get(i.rstrip("ab"))
            if hit and hit[1]:
                expected = hit[1]
                break
        rows.append({
            "id": " / ".join(ids) or "—",
            "module": MOBILE_MODULES[spec],
            "title": title.split(": ", 1)[-1] if ids else title,
            "spec": spec,
            "status": status,
            "expected": expected,
            "note": note,
        })
    return rows


# ------------------------------------------------------------------ Findings
# (ID, area, severity, finding, evidence, status). Web findings from the
# run are appended in main() when a web test fails.
FINDINGS = [
    ("F-01", "Mobile · Library / Home", "Medium",
     "App still offers \"Add a PDF\" (opens the file picker), but the backend refuses every own-PDF upload.",
     "POST /api/documents → 403 \"Uploading your own book is no longer available…\" (rechecked 2026-09-30). "
     "Home's PDF section shows only built-in examples.", "Open (reported 2026-09-24)"),
    ("F-02", "Mobile · Audiobooks", "Medium",
     "\"One narration at a time\" isn't enforced: a second Generate created a second job while the app said "
     "\"A narration is already running\".",
     "2026-09-30 05:11 UTC: job a2b30fa3 created while job ce4375c7 was still running.", "Open"),
    ("F-03", "Backend · Narration queue", "Medium",
     "A short narration sat in \"running\" for ~20 min (progress stuck at 25%, then 100% without finishing); "
     "the app showed it as QUEUED meanwhile.",
     "Job ce4375c7, 2026-09-30 05:09 → done after ~20 min; other jobs that day took under 1 min.", "Observed once"),
    ("F-04", "Mobile · Audiobooks", "Info",
     "First Generate after install shows Android's notification-permission prompt over the \"Narration started\" dialog.",
     "Expected Android 13+ behaviour; tests now tap Allow.", "Not a bug"),
    ("F-05", "Web · Test environment", "High",
     "The admin test account can no longer sign in, blocking 4 admin-only web tests (Home as admin, PDF catalogue, "
     "PDF upload, PDF → narration E2E).",
     "admin account from MOBILE_TEST_ADMIN_EMAIL: UI \"Invalid email or password\"; POST /api/auth/login → 401 (2026-10-02).",
     "Needs new credentials"),
    ("F-06", "Web · Security", "High",
     "Voice-preview authorization (issue #13): an unrelated account could fetch another user's private voice preview.",
     "2026-10-02: a brand-new account requesting the test account's HH2 preview gets HTTP 403 — fixed.",
     "Fixed (verified)"),
    ("F-07", "Web · Sign-in / sign-up", "Info",
     "New /welcome step after sign-in and sign-up: \"What brings you here?\" (Listen / Preserve a voice / Feel close "
     "to someone) with Continue / Skip for now.",
     "Product change; it broke every web test that waited for Home. Tests now skip it.", "Not a bug"),
    ("F-08", "Web · Profile", "Info",
     "Delete account now also asks for the account password (plus optional reason and the DELETE confirmation).",
     "Security improvement; TC-012 updated.", "Not a bug"),
]

# ------------------------------------------------------------------ Workbook
HEAD = PatternFill("solid", fgColor="1F3A5F")
thin = Side(style="thin", color="C9CED6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS_FILL = {PASS: "C6EFCE", FAIL: "FFC7CE", SKIP: "FFEB9C", BLOCKED: "D9D9D9"}
SEV_FILL = {"High": "FDE2E1", "Medium": "FFF4D6", "Low": "E6F4EA", "Info": "EAF1FB"}


def table(ws, cols, widths, rows, status_col=None):
    ws.append(cols)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 28
    for r in rows:
        ws.append(r)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        if status_col:
            st = ws.cell(row=r, column=status_col)
            if st.value in STATUS_FILL:
                st.fill = PatternFill("solid", fgColor=STATUS_FILL[st.value])
                st.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def main():
    reports = sys.argv[1:] or [os.path.join(os.environ.get("TEMP", "."), "odiobuk-web-results.json")]
    web = web_rows(reports)
    mobile = mobile_rows()
    findings = list(FINDINGS)
    for r in web:
        if r["status"] == FAIL:
            findings.append((f"W-{r['id']}", f"Web · {r['module']}", "High" if r["id"].startswith("SEC") else "Medium",
                             f"{r['id']} failed: {r['title']}", r["note"], "Open"))

    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    s["A1"] = "Odiobuk / MemoryWave — Test Report"
    s["A1"].font = Font(bold=True, size=14)
    s["A2"] = f"Web: full cycle {WEB_RUN_DATE} · {WEB_ENV}"
    s["A3"] = f"Mobile: full cycle {MOBILE_RUN_DATE} · {MOBILE_ENV}"
    s["A4"] = ("Pass = behaved as expected · Fail = app/backend defect · Blocked = could not run (environment, e.g. "
               "test account) · Skipped = needs data the backend no longer allows · reasons in the Notes column")
    s.append([])
    cols = ["Platform / Module", "Total", "Pass", "Fail", "Blocked", "Skipped", "Pass % (of run)"]
    s.append(cols)
    hdr = s.max_row
    for c in range(1, len(cols) + 1):
        s.cell(row=hdr, column=c).font = Font(bold=True, color="FFFFFF")
        s.cell(row=hdr, column=c).fill = HEAD

    def line(label, rows, bold=False):
        n = len(rows)
        p = sum(r["status"] == PASS for r in rows)
        f = sum(r["status"] == FAIL for r in rows)
        b = sum(r["status"] == BLOCKED for r in rows)
        k = sum(r["status"] == SKIP for r in rows)
        run = p + f
        s.append([label, n, p, f, b, k, f"{round(100 * p / run)}%" if run else "—"])
        if bold:
            for c in range(1, len(cols) + 1):
                s.cell(row=s.max_row, column=c).font = Font(bold=True)

    for platform, rows in (("WEB", web), ("MOBILE", mobile)):
        line(platform, rows, bold=True)
        for m in dict.fromkeys(r["module"] for r in rows):
            line(f"   {m}", [r for r in rows if r["module"] == m])
    line("TOTAL (web + mobile)", web + mobile, bold=True)
    for r in range(hdr + 1, s.max_row + 1):
        s.cell(row=r, column=3).fill = PatternFill("solid", fgColor="E6F4EA")
        s.cell(row=r, column=4).fill = PatternFill("solid", fgColor="FDE2E1")
    s.append([])
    s.append([f"Findings: {len(findings)} (see the Findings sheet)"])
    s.cell(row=s.max_row, column=1).font = Font(bold=True)
    s.column_dimensions["A"].width = 42
    for col in "BCDEFG":
        s.column_dimensions[col].width = 11

    w = wb.create_sheet("Web Results")
    table(w, ["TC ID", "Module", "Test", "Type", "Spec file", "Status", "Duration (s)", "Actual Result / Notes"],
          [10, 22, 60, 13, 30, 10, 12, 60],
          [[r["id"], r["module"], r["title"], r["type"], r["spec"], r["status"], r["seconds"], r["note"]] for r in web],
          status_col=6)

    m = wb.create_sheet("Mobile Results")
    table(m, ["TC ID(s)", "Module", "Test", "Expected Result (TC matrix)", "Spec file", "Status", "Actual Result / Notes"],
          [16, 18, 52, 56, 26, 10, 52],
          [[r["id"], r["module"], r["title"], r["expected"], r["spec"], r["status"], r["note"]] for r in mobile],
          status_col=6)

    f = wb.create_sheet("Findings")
    table(f, ["ID", "Area", "Severity", "Finding", "Evidence", "Status"], [8, 24, 10, 60, 60, 18], [list(x) for x in findings])
    for r in range(2, f.max_row + 1):
        sev = f.cell(row=r, column=3)
        if sev.value in SEV_FILL:
            sev.fill = PatternFill("solid", fgColor=SEV_FILL[sev.value])
            sev.font = Font(bold=True)

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Odiobuk_Test_Report.xlsx")
    wb.save(out)
    count = lambda rows, st: sum(r["status"] == st for r in rows)
    print(out, f"| web {len(web)}: {count(web, PASS)} pass {count(web, FAIL)} fail {count(web, BLOCKED)} blocked {count(web, SKIP)} skip",
          f"| mobile {len(mobile)}: {count(mobile, PASS)} pass {count(mobile, SKIP)} skip | findings {len(findings)}")


if __name__ == "__main__":
    main()
