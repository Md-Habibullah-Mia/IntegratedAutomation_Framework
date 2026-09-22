import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DOCS = Path(__file__).parent
SOURCES = [
    ("Onboarding & Auth", DOCS / "tc_auth.tsv"),
    ("Home & Library", DOCS / "tc_home_library.tsv"),
    ("Audiobooks & Player", DOCS / "tc_audiobooks.tsv"),
    ("Voices & Recording", DOCS / "tc_voices.tsv"),
    ("Profile & Vault", DOCS / "tc_profile_vault.tsv"),
]
OUT_PATH = DOCS / "MemoryWave_Test_Cases.xlsx"

HEADERS = ["TC ID", "Title", "Preconditions", "Steps", "Expected Result", "Actual Result"]

def parse_tsv(path: Path):
    rows = []
    with path.open("r", encoding="utf-8") as f:
        reader = csv.reader(f, delimiter="|")
        first = True
        for raw in reader:
            if not raw or not any(c.strip() for c in raw):
                continue
            cells = [c.strip() for c in raw]
            if first:
                first = False
                continue  # skip header row
            if len(cells) < 5:
                continue
            tc_id, title, pre, steps, expected = cells[0], cells[1], cells[2], cells[3], cells[4]
            rows.append((tc_id, title, pre, steps, expected))
    return rows

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)
title_font = Font(bold=True, size=14)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wrap_center = Alignment(wrap_text=True, vertical="top", horizontal="center")

all_rows_for_summary = []
grand_total = 0

for sheet_name, path in SOURCES:
    rows = parse_tsv(path)
    grand_total += len(rows)
    all_rows_for_summary.append((sheet_name, len(rows)))

    ws = wb.create_sheet(title=sheet_name[:31])
    ws.append(HEADERS)
    for col_idx, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    for r in rows:
        ws.append(list(r) + [""])  # blank Actual Result

    n = len(rows)
    for row_idx in range(2, n + 2):
        for col_idx in range(1, 7):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = wrap_center if col_idx == 1 else wrap

    widths = {"A": 12, "B": 34, "C": 34, "D": 60, "E": 60, "F": 22}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    if n > 0:
        table_ref = f"A1:F{n+1}"
        table = Table(displayName=f"T_{re.sub(r'[^A-Za-z0-9]', '', sheet_name)}", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)

# Summary sheet first
summary = wb.create_sheet(title="Summary", index=0)
summary["A1"] = "MemoryWave (Odiobuk Mobile App) — Test Case Matrix"
summary["A1"].font = title_font
summary["A2"] = "Derived from the Flutter source (D:\\My Projects\\Audiobook-mobile) and live app exploration."
summary["A2"].font = Font(italic=True, size=10, color="555555")
summary["A4"] = "Area"
summary["B4"] = "Test Cases"
summary["A4"].font = header_font
summary["B4"].font = header_font
summary["A4"].fill = header_fill
summary["B4"].fill = header_fill
row = 5
for name, count in all_rows_for_summary:
    summary.cell(row=row, column=1, value=name)
    summary.cell(row=row, column=2, value=count)
    row += 1
summary.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
summary.cell(row=row, column=2, value=grand_total).font = Font(bold=True)
summary.column_dimensions["A"].width = 30
summary.column_dimensions["B"].width = 14

wb.save(OUT_PATH)
print(f"Wrote {OUT_PATH} with {grand_total} test cases across {len(SOURCES)} sheets")
